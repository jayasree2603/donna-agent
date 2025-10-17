import re
import os
import time
import asyncio
import logging
import sys
import uuid
import datetime
import json
from pathlib import Path
from docx import Document
from dotenv import load_dotenv
from typing import Dict, List, Optional, Tuple

from google.adk.agents import Agent
from google.adk.sessions import InMemorySessionService
from google.adk.tools import FunctionTool
from google.adk.models.lite_llm import LiteLlm
from google.adk.runners import Runner
from google.genai import types


load_dotenv()

# ============================================================================
# CONFIGURATION
# ============================================================================

MAX_FILENAME_LENGTH = 200
MAX_VARIABLE_VALUE_LENGTH = 10000
MAX_VARIABLES_COUNT = 100
ALLOWED_EXTENSIONS = {'.docx'}
MAX_FILE_SIZE_MB = 50
FORBIDDEN_CHARS = ['<', '>', ':', '"', '/', '\\', '|', '?', '*']

TEMPLATES_DIR = os.getenv("TEMPLATES_DIR", "./templates")
RESULTS_DIR = os.getenv("RESULTS_DIR", "./results")

# ============================================================================
# VALIDATION FUNCTIONS
# ============================================================================

def validate_email(email: str) -> Tuple[bool, Optional[str]]:
    """Validate email format."""
    email_pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    if not re.match(email_pattern, email):
        return False, "Invalid email format. Please use format: example@domain.com"
    return True, None


def validate_phone(phone: str) -> Tuple[bool, Optional[str]]:
    """Validate phone number format."""
    cleaned = re.sub(r'[\s\-\(\)\.+]', '', phone)
    if not cleaned.isdigit() or len(cleaned) < 7 or len(cleaned) > 15:
        return False, "Invalid phone number. Please use a valid phone format (7-15 digits)"
    return True, None


def validate_and_normalize_date(date_str: str) -> Tuple[bool, Optional[str], Optional[str]]:
    """Validate and normalize date."""
    date_str = date_str.strip()
    
    accepted_formats = [
        '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%m-%d-%Y',
        '%d/%m/%y', '%m/%d/%y', '%Y-%m-%d',
    ]
    
    for fmt in accepted_formats:
        try:
            parsed_date = datetime.datetime.strptime(date_str, fmt)
            normalized = parsed_date.strftime('%d/%m/%Y')
            return True, None, normalized
        except ValueError:
            continue
    
    return False, "Invalid date. Please use format: DD/MM/YYYY", None


def validate_input_by_type(variable_name: str, value: str) -> Tuple[bool, Optional[str], Optional[str]]:
    """Validate input based on variable name pattern."""
    if not value or not value.strip():
        return False, f"{variable_name} cannot be empty", None
    
    value = value.strip()
    
    if 'email' in variable_name.lower():
        is_valid, error = validate_email(value)
        return is_valid, error, value if is_valid else None
    
    if 'phone' in variable_name.lower():
        is_valid, error = validate_phone(value)
        return is_valid, error, value if is_valid else None
    
    if 'date' in variable_name.lower() and 'candidate' not in variable_name.lower():
        return validate_and_normalize_date(value)
    
    if 'name' in variable_name.lower():
        if not any(term in variable_name.lower() for term in ['filename', 'username', 'hostname']):
            if not re.match(r"^[a-zA-Z\s\-']+$", value):
                return False, f"{variable_name} should only contain letters, spaces, hyphens, or apostrophes", None
            if len(value) < 2:
                return False, f"{variable_name} should be at least 2 characters long", None
    
    return True, None, value


def sanitize_filename(filename: str) -> str:
    """Sanitize filename to remove dangerous characters."""
    for char in FORBIDDEN_CHARS:
        filename = filename.replace(char, '_')
    filename = ''.join(char for char in filename if ord(char) >= 32)
    filename = filename.strip()
    if not filename:
        filename = "document"
    if len(filename) > MAX_FILENAME_LENGTH:
        filename = filename[:MAX_FILENAME_LENGTH]
    return filename


# ============================================================================
# TOOL FUNCTIONS
# ============================================================================

def list_available_templates(directory_path: str = TEMPLATES_DIR) -> dict:
    """
    List all available document templates.
    
    Args:
        directory_path: Path to the templates directory
        
    Returns:
        Dictionary with list of available templates or error message
    """
    if not os.path.isdir(directory_path):
        return {
            "success": False,
            "error": f"Template directory not found: {directory_path}",
            "message": "Template directory not found. Please ensure './templates' exists."
        }
    
    try:
        templates = []
        for filename in os.listdir(directory_path):
            if filename.endswith("_template.docx"):
                friendly_name = filename.replace("_template.docx", "").replace("_", " ")
                templates.append({
                    "name": friendly_name,
                    "filename": filename,
                    "path": os.path.join(directory_path, filename)
                })
        
        if not templates:
            return {
                "success": False,
                "error": "No templates found",
                "message": f"No templates found in {directory_path}. Please add *_template.docx files."
            }
        
        return {
            "success": True,
            "templates": templates,
            "count": len(templates),
            "message": f"Found {len(templates)} available templates."
        }
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "message": f"Error reading templates directory: {str(e)}"
        }


def find_template(user_query: str, directory_path: str = TEMPLATES_DIR) -> dict:
    """
    Find the template that matches the user's query.
    
    Args:
        user_query: The user's request describing what template they want
        directory_path: Path to the directory containing templates
        
    Returns:
        Dictionary with template information or error message
    """
    if not os.path.isdir(directory_path):
        return {
            "success": False,
            "error": f"Directory not found: {directory_path}",
            "message": "Template directory not found."
        }
    
    templates = {}
    try:
        for filename in os.listdir(directory_path):
            if filename.endswith("_template.docx"):
                friendly_name = filename.replace("_template.docx", "").replace("_", " ")
                templates[friendly_name.lower()] = {
                    "name": friendly_name,
                    "path": os.path.join(directory_path, filename)
                }
    except Exception as e:
        return {
            "success": False,
            "error": f"Could not read directory: {e}",
            "message": "Error reading templates directory."
        }
    
    if not templates:
        return {
            "success": False,
            "error": "No templates found",
            "message": "No templates available.",
            "available_templates": []
        }
    
    # Try to find a direct match
    query_lower = user_query.lower()
    if query_lower in templates:
        return {
            "success": True,
            "template_name": templates[query_lower]["name"],
            "template_path": templates[query_lower]["path"],
            "status": "exact_match",
            "message": f"Found template: {templates[query_lower]['name']}"
        }
    
    # Try to find a partial match
    for template_key, template_info in templates.items():
        if query_lower in template_key or template_key in query_lower:
            return {
                "success": True,
                "template_name": template_info["name"],
                "template_path": template_info["path"],
                "status": "partial_match",
                "message": f"Found matching template: {template_info['name']}"
            }
    
    # No match found, return available templates
    return {
        "success": False,
        "status": "no_match",
        "available_templates": [t["name"] for t in templates.values()],
        "message": "No matching template found. Please choose from the available templates.",
        "error": "No matching template found"
    }


def get_template_variables(template_path: str) -> dict:
    """
    Extract variables from a template.
    
    Args:
        template_path: Path to the template file
        
    Returns:
        Dictionary with variables found in the template
    """
    try:
        if not os.path.exists(template_path):
            return {
                "success": False,
                "error": f"Template file not found: {template_path}",
                "message": "Template file not found."
            }
        
        doc = Document(template_path)
        full_text = []
        
        for para in doc.paragraphs:
            full_text.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        full_text.append(para.text)
        
        placeholders = re.findall(r'\{(\w+)\}', "\n".join(full_text))
        unique_placeholders = sorted(list(set(placeholders)))
        
        if not unique_placeholders:
            return {
                "success": False,
                "status": "no_variables",
                "message": "No variables found in the template.",
                "variables": []
            }
        
        return {
            "success": True,
            "status": "variables_found",
            "variables": unique_placeholders,
            "count": len(unique_placeholders),
            "message": f"Found {len(unique_placeholders)} variables: {', '.join(unique_placeholders)}"
        }
        
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "message": f"Error reading template file: {str(e)}"
        }


def validate_variable_value(variable_name: str, value: str) -> dict:
    """
    Validate a variable value based on its name.
    
    Args:
        variable_name: The name of the variable being validated
        value: The value to validate
        
    Returns:
        Dictionary with validation result
    """
    is_valid, error_msg, normalized_value = validate_input_by_type(variable_name, value)
    
    return {
        "success": is_valid,
        "variable_name": variable_name,
        "original_value": value,
        "normalized_value": normalized_value if is_valid else value,
        "error": error_msg if not is_valid else None,
        "message": "Valid" if is_valid else error_msg
    }


def generate_document(template_path: str, variables_json: str) -> dict:
    """
    Fill the template with provided variables and generate the document.
    
    Args:
        template_path: Path to the template file
        variables_json: JSON string containing variable names and their values
        
    Returns:
        Dictionary with status and path to the generated document
    """
    try:
        # Parse variables
        variables = json.loads(variables_json)
        
        if not os.path.exists(template_path):
            return {
                "success": False,
                "error": "Template file not found",
                "message": "Template file not found."
            }
        
        doc = Document(template_path)
        replacements = {f"{{{key}}}": str(value) for key, value in variables.items()}
        replacements_made = 0
        
        def replace_in_paragraph(paragraph, replacements):
            nonlocal replacements_made
            full_text = "".join(run.text for run in paragraph.runs)
            
            has_replacement = False
            for placeholder, value in replacements.items():
                if placeholder in full_text:
                    has_replacement = True
                    full_text = full_text.replace(placeholder, value)
                    replacements_made += 1
            
            if has_replacement:
                for run in paragraph.runs:
                    run.text = ""
                if paragraph.runs:
                    paragraph.runs[0].text = full_text
                else:
                    paragraph.add_run(full_text)
        
        # Replace in paragraphs
        for para in doc.paragraphs:
            replace_in_paragraph(para, replacements)
        
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        replace_in_paragraph(para, replacements)
        
        # Ensure results directory exists
        os.makedirs(RESULTS_DIR, exist_ok=True)
        
        # Generate unique filename
        base_filename = os.path.basename(template_path).replace("_template.docx", "")
        base_filename = sanitize_filename(base_filename)
        
        timestamp = int(time.time())
        unique_id = uuid.uuid4().hex[:8]
        new_filename = f"{base_filename}_{timestamp}_{unique_id}.docx"
        new_filepath = os.path.join(RESULTS_DIR, new_filename)
        
        doc.save(new_filepath)
        
        if not os.path.exists(new_filepath):
            return {
                "success": False,
                "error": "File not saved successfully",
                "message": "Document was not saved successfully."
            }
        
        return {
            "success": True,
            "status": "success",
            "file_path": new_filepath,
            "filename": new_filename,
            "replacements_made": replacements_made,
            "message": f"✓ Document created successfully! File saved at: {new_filepath}"
        }
        
    except json.JSONDecodeError as e:
        return {
            "success": False,
            "error": f"Invalid JSON format: {e}",
            "message": "Invalid variable format provided."
        }
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "message": f"Error generating document: {str(e)}"
        }


# ============================================================================
# TOOL DEFINITIONS
# ============================================================================

list_templates_tool = FunctionTool(list_available_templates)
find_template_tool = FunctionTool(find_template)
get_variables_tool = FunctionTool(get_template_variables)
validate_value_tool = FunctionTool(validate_variable_value)
generate_doc_tool = FunctionTool(generate_document)

# ============================================================================
# AGENT SETUP
# ============================================================================

model = LiteLlm(model="gpt-4o-mini")

instruction = """You are Donna, a friendly and efficient HR Document Assistant.

Your role is to help users create HR documents through a conversational workflow:

1. **Understand Intent**: When a user requests a document, use find_template to locate the appropriate template from ./templates directory.

2. **List Templates**: If unsure or if no match is found, use list_available_templates to show options.

3. **Extract Variables**: Once a template is found, use get_template_variables to identify what information is needed.

4. **Collect Information**: Ask the user for each variable ONE AT A TIME in a natural, conversational way.
   - Before accepting any value, ALWAYS use validate_variable_value to check if it's valid
   - If validation fails, explain the error clearly and ask again
   - Keep track of what you've collected and what's still needed

5. **Generate Document**: Once ALL variables are collected and validated, use generate_document with the template path and a JSON string of all variables.

6. **Confirm Success**: Present the file path clearly to the user.

IMPORTANT GUIDELINES:
- Be conversational and friendly, not robotic
- Ask for ONE piece of information at a time
- ALWAYS validate inputs before accepting them
- If validation fails, explain why and ask again
- Keep track of progress (e.g., "Great! 3 down, 2 to go")
- Use the tools provided - don't make up information
- All templates are in './templates' directory
- All generated documents go to './results' directory

Example flow:
User: "I need an offer letter"
You: Use find_template("offer letter", "./templates")
Then: Use get_template_variables(template_path)
Then: "Great! I found the offer letter template. I need to collect 5 pieces of information. Let's start - what is the candidate_name?"
User: "John Doe"
You: Use validate_variable_value("candidate_name", "John Doe")
Then: "Perfect! What is the position?"
...and so on until all variables are collected.
Finally: Use generate_document(template_path, json_variables)
"""

root_agent = Agent(
    name="Donna",
    model=model,
    tools=[
        list_templates_tool,
        find_template_tool,
        get_variables_tool,
        validate_value_tool,
        generate_doc_tool
    ],
    instruction=instruction
)

# ============================================================================
# SESSION AND RUNNER SETUP
# ============================================================================

APP_NAME = "Donna-HR-Assistant"
USER_ID = "local-user"

session_service = InMemorySessionService()

runner = Runner(
    agent=root_agent,
    app_name=APP_NAME,
    session_service=session_service
)

# ============================================================================
# MAIN ASYNC LOOP
# ============================================================================

async def run_donna_assistant():
    """Main async loop for the Donna assistant."""
    print("=" * 60)
    print("Donna - HR Document Assistant (Google ADK)")
    print("=" * 60)
    print("Type 'quit', 'exit', or 'bye' to end the session.")
    print("Type 'reset' to start a new conversation.\n")
    
    # Create session
    session = await session_service.create_session(
        app_name=APP_NAME,
        user_id=USER_ID,
        state={}
    )
    session_id = session.id
    
    print("Agent: Hello! I'm Donna, your HR Document Assistant.")
    print("       I can help you create various HR documents.")
    print("       What document would you like to create today?\n")
    
    while True:
        try:
            user_input = await asyncio.to_thread(input, "You: ")
        except (EOFError, KeyboardInterrupt):
            print("\n\nGoodbye! Have a great day!")
            break
        
        if not user_input.strip():
            continue
        
        if user_input.lower() in ['quit', 'exit', 'bye']:
            print("\nAgent: Goodbye! Have a great day!")
            break
        
        if user_input.lower() == 'reset':
            # Create new session
            session = await session_service.create_session(
                app_name=APP_NAME,
                user_id=USER_ID,
                state={}
            )
            session_id = session.id
            print("\nAgent: Conversation reset. What document would you like to create?\n")
            continue
        
        print("Agent: ", end="", flush=True)
        
        # Create content message
        content = types.Content(
            role='user',
            parts=[types.Part(text=user_input)]
        )
        
        # Stream the response
        try:
            for event in runner.run(
                user_id=USER_ID,
                session_id=session_id,
                new_message=content
            ):
                if event.content and event.content.parts:
                    for part in event.content.parts:
                        if hasattr(part, 'text') and part.text:
                            print(part.text, end="", flush=True)
            print()  # Newline after response
        except KeyboardInterrupt:
            print("\n\nInterrupted. Goodbye!")
            break
        except Exception as e:
            print(f"\nAn error occurred: {str(e)}")
            print("Let's try again. What would you like to do?\n")


# ============================================================================
# MAIN EXECUTION
# ============================================================================

if __name__ == "__main__":
    # Suppress logging
    logging.getLogger("litellm").setLevel(logging.CRITICAL)
    logging.getLogger("litellm.logging_utils").setLevel(logging.CRITICAL)
    
    # Redirect stderr to suppress error messages
    original_stderr = sys.stderr
    sys.stderr = open(os.devnull, 'w')
    
    try:
        asyncio.run(run_donna_assistant())
    finally:
        sys.stderr = original_stderr