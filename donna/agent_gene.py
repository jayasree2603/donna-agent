import re
import os
import time
import asyncio
import logging
import sys
import uuid
import datetime
import json
from pathlib import Path
from docx import Document
from dotenv import load_dotenv
from typing import Dict, List, Optional, Tuple

from google.adk.agents import Agent
from google.adk.sessions import InMemorySessionService
from google.adk.tools import FunctionTool
from google.adk.models.lite_llm import LiteLlm
from google.adk.runners import Runner
from google.genai import types
from .config import TEMPLATES_DIR, RESULTS_DIR, MAX_FILENAME_LENGTH, FORBIDDEN_CHARS, ALLOWED_EXTENSIONS

load_dotenv()

# ============================================================================
# VALIDATION FUNCTIONS
# ============================================================================

def validate_email(email: str) -> Tuple[bool, Optional[str]]:
    email_pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    if not re.match(email_pattern, email):
        return False, "Invalid email format. Please use format: example@domain.com"
    return True, None


def validate_phone(phone: str) -> Tuple[bool, Optional[str]]:
    cleaned = re.sub(r'[\s\-\(\)\.+]', '', phone)
    if not cleaned.isdigit() or len(cleaned) < 7 or len(cleaned) > 15:
        return False, "Invalid phone number. Please use a valid phone format (7-15 digits)"
    return True, None


def validate_and_normalize_date(date_str: str) -> Tuple[bool, Optional[str], Optional[str]]:
    date_str = date_str.strip()
    accepted_formats = [
        '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%m-%d-%Y',
        '%d/%m/%y', '%m/%d/%y', '%Y-%m-%d',
    ]
    for fmt in accepted_formats:
        try:
            parsed_date = datetime.datetime.strptime(date_str, fmt)
            return True, None, parsed_date.strftime('%d/%m/%Y')
        except ValueError:
            continue
    return False, "Invalid date. Please use format: DD/MM/YYYY", None


def validate_input_by_type(variable_name: str, value: str) -> Tuple[bool, Optional[str], Optional[str]]:
    if not value or not value.strip():
        return False, f"{variable_name} cannot be empty", None
    value = value.strip()
    if 'email' in variable_name.lower():
        is_valid, error = validate_email(value)
        return is_valid, error, value if is_valid else None
    if 'phone' in variable_name.lower():
        is_valid, error = validate_phone(value)
        return is_valid, error, value if is_valid else None
    if 'date' in variable_name.lower() and 'candidate' not in variable_name.lower():
        return validate_and_normalize_date(value)
    if 'name' in variable_name.lower():
        if not any(term in variable_name.lower() for term in ['filename', 'username', 'hostname']):
            if not re.match(r"^[a-zA-Z\s\-']+$", value):
                return False, f"{variable_name} should only contain letters, spaces, hyphens, or apostrophes", None
            if len(value) < 2:
                return False, f"{variable_name} should be at least 2 characters long", None
    return True, None, value


def sanitize_filename(filename: str) -> str:
    for char in FORBIDDEN_CHARS:
        filename = filename.replace(char, '_')
    filename = ''.join(char for char in filename if ord(char) >= 32)
    filename = filename.strip()
    if not filename:
        filename = "document"
    if len(filename) > MAX_FILENAME_LENGTH:
        filename = filename[:MAX_FILENAME_LENGTH]
    return filename


# ============================================================================
# TOOL FUNCTIONS
# ============================================================================

def list_available_templates(directory_path: str = TEMPLATES_DIR) -> dict:
    if not os.path.isdir(directory_path):
        return {"success": False, "error": f"Template directory not found: {directory_path}", "message": "Template directory not found."}
    try:
        templates = []
        for filename in os.listdir(directory_path):
            if filename.endswith("_template.docx"):
                friendly_name = filename.replace("_template.docx", "").replace("_", " ")
                templates.append({"name": friendly_name, "filename": filename, "path": os.path.join(directory_path, filename)})
        if not templates:
            return {"success": False, "error": "No templates found", "message": f"No templates found in {directory_path}."}
        return {"success": True, "templates": templates, "count": len(templates), "message": f"Found {len(templates)} templates."}
    except Exception as e:
        return {"success": False, "error": str(e), "message": f"Error reading templates directory: {str(e)}"}


def find_template(user_query: str, directory_path: str = TEMPLATES_DIR) -> dict:
    # Ensure fallback if empty string is passed
    if not directory_path:
        directory_path = TEMPLATES_DIR
    
    if not os.path.isdir(directory_path):
        return {
            "success": False,
            "error": f"Directory not found: {directory_path}",
            "message": "Template directory not found."
        }

    templates = {}
    try:
        for filename in os.listdir(directory_path):
            if filename.endswith("_template.docx"):
                friendly_name = filename.replace("_template.docx", "").replace("_", " ")
                templates[friendly_name.lower()] = {
                    "name": friendly_name,
                    "path": os.path.join(directory_path, filename)
                }
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "message": "Error reading templates directory."
        }

    if not templates:
        return {
            "success": False,
            "error": "No templates found",
            "available_templates": []
        }

    query_lower = user_query.lower()

    # Exact match
    if query_lower in templates:
        return {
            "success": True,
            "template_name": templates[query_lower]["name"],
            "template_path": templates[query_lower]["path"],
            "status": "exact_match",
            "message": f"Found template: {templates[query_lower]['name']}"
        }

    # Partial match
    for template_key, template_info in templates.items():
        if query_lower in template_key or template_key in query_lower:
            return {
                "success": True,
                "template_name": template_info["name"],
                "template_path": template_info["path"],
                "status": "partial_match",
                "message": f"Found matching template: {template_info['name']}"
            }

    # No match found
    return {
        "success": False,
        "status": "no_match",
        "available_templates": [t["name"] for t in templates.values()],
        "message": "No matching template found.",
        "error": "No matching template found"
    }


def get_template_variables(template_path: str) -> dict:
    if not os.path.exists(template_path):
        return {"success": False, "error": f"Template file not found: {template_path}", "message": "Template file not found."}
    try:
        doc = Document(template_path)
        full_text = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.extend([para.text for para in cell.paragraphs])
        placeholders = re.findall(r'\{(\w+)\}', "\n".join(full_text))
        unique_placeholders = sorted(list(set(placeholders)))
        if not unique_placeholders:
            return {"success": False, "status": "no_variables", "message": "No variables found in the template.", "variables": []}
        return {"success": True, "status": "variables_found", "variables": unique_placeholders, "count": len(unique_placeholders), "message": f"Found {len(unique_placeholders)} variables: {', '.join(unique_placeholders)}"}
    except Exception as e:
        return {"success": False, "error": str(e), "message": f"Error reading template file: {str(e)}"}


def validate_variable_value(variable_name: str, value: str) -> dict:
    is_valid, error_msg, normalized_value = validate_input_by_type(variable_name, value)
    return {"success": is_valid, "variable_name": variable_name, "original_value": value, "normalized_value": normalized_value if is_valid else value, "error": error_msg if not is_valid else None, "message": "Valid" if is_valid else error_msg}


def generate_document(template_path: str, variables_json: str) -> dict:
    try:
        variables = json.loads(variables_json)
        if not os.path.exists(template_path):
            return {"success": False, "error": "Template file not found", "message": "Template file not found."}
        doc = Document(template_path)
        replacements = {f"{{{key}}}": str(value) for key, value in variables.items()}
        replacements_made = 0
        def replace_in_paragraph(paragraph, replacements):
            nonlocal replacements_made
            full_text = "".join(run.text for run in paragraph.runs)
            has_replacement = False
            for placeholder, value in replacements.items():
                if placeholder in full_text:
                    has_replacement = True
                    full_text = full_text.replace(placeholder, value)
                    replacements_made += 1
            if has_replacement:
                for run in paragraph.runs:
                    run.text = ""
                if paragraph.runs:
                    paragraph.runs[0].text = full_text
                else:
                    paragraph.add_run(full_text)
        for para in doc.paragraphs:
            replace_in_paragraph(para, replacements)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        replace_in_paragraph(para, replacements)
        os.makedirs(RESULTS_DIR, exist_ok=True)
        base_filename = sanitize_filename(os.path.basename(template_path).replace("_template.docx", ""))
        timestamp = int(time.time())
        unique_id = uuid.uuid4().hex[:8]
        new_filename = f"{base_filename}_{timestamp}_{unique_id}.docx"
        new_filepath = os.path.join(RESULTS_DIR, new_filename)
        doc.save(new_filepath)
        if not os.path.exists(new_filepath):
            return {"success": False, "error": "File not saved successfully", "message": "Document was not saved successfully."}
        return {"success": True, "status": "success", "file_path": new_filepath, "filename": new_filename, "replacements_made": replacements_made, "message": f"✓ Document created successfully! File saved at: {new_filepath}"}
    except json.JSONDecodeError as e:
        return {"success": False, "error": f"Invalid JSON format: {e}", "message": "Invalid variable format provided."}
    except Exception as e:
        return {"success": False, "error": str(e), "message": f"Error generating document: {str(e)}"}


# ============================================================================
# TOOL DEFINITIONS
# ============================================================================

list_templates_tool = FunctionTool(list_available_templates)
find_template_tool = FunctionTool(find_template)
get_variables_tool = FunctionTool(get_template_variables)
validate_value_tool = FunctionTool(validate_variable_value)
generate_doc_tool = FunctionTool(generate_document)
# ============================================================================
# AGENT SETUP
# ============================================================================

model = LiteLlm(model="gpt-4o-mini")

instruction = instruction = """
You are Donna, an expert HR document specialist. You create HR documents quickly, accurately, and professionally.

RULES:
- Focus exclusively on HR documents, templates, and the information needed.
- Never mention technical operations, backend processes, or system functions.
- Keep communication natural, concise, and professional.
- Ask clarifying questions only if information is genuinely missing.
- Do not assume or guess any folder names.
- Only look for templates in the configured directory.

ESSENCE:
- Anticipate user needs and handle details flawlessly.
- Confident, direct, and approachable.
- Maintain high standards of quality, formatting, and accuracy.
- Celebrate wins with warm, professional responses.

WORKFLOW:
1. Understand what document the user needs.
2. Find the appropriate template automatically.
3. Collect required information intelligently from the user.
4. Validate, format, and fill the template.
5. Confirm completion and provide the document.

COMMUNICATION STYLE:
- Brief, confident introductions.
- Redirect politely if asked about anything outside HR documents.
- Only discuss documents, templates, and required information.
- Handle missing info or obstacles smoothly and suggest alternatives.

DECISION-MAKING:
- Act automatically when requirements are clear.
- Chain tasks seamlessly, using outputs from one step in the next.
- Never ask for permission to proceed.
- Make smart decisions and keep the workflow efficient.

THE DONNA STANDARD:
- Never guess; always use actual information.
- Process input intelligently and strategically.
- Maintain professionalism, clarity, and efficiency at all times.

"""

root_agent = Agent(
    name="Donna",
    model=model,
    tools=[
        list_templates_tool,
        find_template_tool,
        get_variables_tool,
        validate_value_tool,
        generate_doc_tool
    ],
    instruction=instruction
)

# ============================================================================
# SESSION AND RUNNER SETUP
# ============================================================================

APP_NAME = "Donna-HR-Assistant"
USER_ID = "local-user"

session_service = InMemorySessionService()

runner = Runner(
    agent=root_agent,
    app_name=APP_NAME,
    session_service=session_service
)

# ============================================================================
# MAIN ASYNC LOOP
# ============================================================================

async def run_donna_assistant():
    """Main async loop for the Donna assistant."""
    print("=" * 60)
    print("Donna - HR Document Assistant (Google ADK)")
    print("=" * 60)
    print("Type 'quit', 'exit', or 'bye' to end the session.")
    print("Type 'reset' to start a new conversation.\n")
    
    # Create session
    session = await session_service.create_session(
        app_name=APP_NAME,
        user_id=USER_ID,
        state={}
    )
    session_id = session.id
    
    print("Agent: Hello! I'm Donna, your HR Document Assistant.")
    print("       I can help you create various HR documents.")
    print("       What document would you like to create today?\n")
    
    while True:
        try:
            user_input = await asyncio.to_thread(input, "You: ")
        except (EOFError, KeyboardInterrupt):
            print("\n\nGoodbye! Have a great day!")
            break
        
        if not user_input.strip():
            continue
        
        if user_input.lower() in ['quit', 'exit', 'bye']:
            print("\nAgent: Goodbye! Have a great day!")
            break
        
        if user_input.lower() == 'reset':
            # Create new session
            session = await session_service.create_session(
                app_name=APP_NAME,
                user_id=USER_ID,
                state={}
            )
            session_id = session.id
            print("\nAgent: Conversation reset. What document would you like to create?\n")
            continue
        
        print("Agent: ", end="", flush=True)
        
        # Create content message
        content = types.Content(
            role='user',
            parts=[types.Part(text=user_input)]
        )
        
        # Stream the response
        try:
            for event in runner.run(
                user_id=USER_ID,
                session_id=session_id,
                new_message=content
            ):
                if event.content and event.content.parts:
                    for part in event.content.parts:
                        if hasattr(part, 'text') and part.text:
                            print(part.text, end="", flush=True)
            print()  # Newline after response
        except KeyboardInterrupt:
            print("\n\nInterrupted. Goodbye!")
            break
        except Exception as e:
            print(f"\nAn error occurred: {str(e)}")
            print("Let's try again. What would you like to do?\n")


# ============================================================================
# MAIN EXECUTION
# ============================================================================

if __name__ == "__main__":
    # Suppress logging
    logging.getLogger("litellm").setLevel(logging.CRITICAL)
    logging.getLogger("litellm.logging_utils").setLevel(logging.CRITICAL)
    
    # Redirect stderr to suppress error messages
    original_stderr = sys.stderr
    sys.stderr = open(os.devnull, 'w')
    
    try:
        asyncio.run(run_donna_assistant())
    finally:
        sys.stderr = original_stderr