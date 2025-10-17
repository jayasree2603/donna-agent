import re
import os
import time
import asyncio
import logging
import sys
import uuid
import datetime
import json
from pathlib import Path
from docx import Document
from dotenv import load_dotenv
from typing import Dict, List, Optional, Tuple

# Google ADK imports
from google.adk.agents import Agent
from google.adk.sessions import InMemorySessionService
from google.adk.tools import FunctionTool
from google.adk.models.lite_llm import LiteLlm

load_dotenv()

# ============================================================================
# VALIDATION CONFIGURATION
# ============================================================================


MAX_FILENAME_LENGTH = 200
MAX_VARIABLE_VALUE_LENGTH = 10000
MAX_VARIABLES_COUNT = 100
ALLOWED_EXTENSIONS = {'.docx'}
MAX_FILE_SIZE_MB = 50
FORBIDDEN_CHARS = ['<', '>', ':', '"', '/', '\\', '|', '?']

TEMPLATES_DIR = os.getenv("TEMPLATES_DIR", "./templates")
RESULTS_DIR = os.getenv("RESULTS_DIR", "./results")

# ============================================================================
# INPUT VALIDATION FUNCTIONS
# ============================================================================

def validate_email(email: str) -> Tuple[bool, Optional[str]]:
    """Validate email format."""
    email_pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    if not re.match(email_pattern, email):
        return False, "Invalid email format. Please use format: example@domain.com"
    return True, None


def validate_phone(phone: str) -> Tuple[bool, Optional[str]]:
    """Validate phone number format."""
    # Remove common separators
    cleaned = re.sub(r'[\s\-\(\)\.+]', '', phone)
    
    # Check if it contains only digits and is 7-15 digits long
    if not cleaned.isdigit() or len(cleaned) < 7 or len(cleaned) > 15:
        return False, "Invalid phone number. Please use a valid phone format (7-15 digits)"
    
    return True, None


def validate_date(date_str: str) -> Tuple[bool, Optional[str]]:
    """Validate date format. Accepts multiple formats."""
    date_str = date_str.strip()
    
    # List of accepted date formats
    accepted_formats = [
        '%d/%m/%Y',    # 30/01/2024
        '%m/%d/%Y',    # 01/30/2024
        '%d-%m-%Y',    # 30-01-2024
        '%m-%d-%Y',    # 01-30-2024
        '%d/%m/%y',    # 30/01/24
        '%m/%d/%y',    # 01/30/24
        '%Y-%m-%d',    # 2024-01-30
    ]
    
    for fmt in accepted_formats:
        try:
            datetime.datetime.strptime(date_str, fmt)
            return True, None
        except ValueError:
            continue
    
    return False, f"Invalid date format. Please use one of: DD/MM/YYYY, MM/DD/YYYY, DD-MM-YYYY, or YYYY-MM-DD"


def validate_input_by_type(variable_name: str, value: str) -> Tuple[bool, Optional[str]]:
    """
    Validate input based on variable name pattern.
    
    Returns:
        Tuple of (is_valid, error_message)
    """
    # Check for empty input
    if not value or not value.strip():
        return False, f"{variable_name} cannot be empty"
    
    value = value.strip()
    
    # Email validation (must have 'email' in name)
    if 'email' in variable_name.lower():
        return validate_email(value)
    
    # Phone validation (must have 'phone' in name)
    if 'phone' in variable_name.lower():
        return validate_phone(value)
    
    # Date validation (must have 'date' in name, not just 'candidate')
    if 'date' in variable_name.lower() and 'candidate' not in variable_name.lower():
        return validate_date(value)
    
    # Name validation (only if contains 'name' but NOT other keywords like 'filename', 'username')
    if 'name' in variable_name.lower():
        # Skip validation for compound terms
        if not any(term in variable_name.lower() for term in ['filename', 'username', 'hostname', 'pathname']):
            if not re.match(r"^[a-zA-Z\s\-']+$", value):
                return False, f"{variable_name} should only contain letters, spaces, hyphens, or apostrophes"
            if len(value) < 2:
                return False, f"{variable_name} should be at least 2 characters long"
    
    # Title/Position validation (allow alphanumeric, spaces, hyphens)
    if any(word in variable_name.lower() for word in ['title', 'position', 'department', 'role']):
        if not re.match(r"^[a-zA-Z0-9\s\-/&]+$", value):
            return False, f"{variable_name} contains invalid characters"
    
    return True, None


def validate_template_path(template_path: str) -> Tuple[bool, Optional[str]]:
    """Validate the template file path."""
    if not template_path or not isinstance(template_path, str):
        return False, "Template path must be a non-empty string"
    
    if '..' in template_path:
        return False, "Invalid path: Path traversal detected"
    
    try:
        path = Path(template_path)
    except Exception as e:
        return False, f"Invalid path format: {e}"
    
    if not path.exists():
        return False, f"Template file not found: {template_path}"
    
    if not path.is_file():
        return False, f"Path is not a file: {template_path}"
    
    if path.suffix.lower() not in ALLOWED_EXTENSIONS:
        return False, f"Invalid file type. Expected .docx, got: {path.suffix}"
    
    file_size_mb = path.stat().st_size / (1024 * 1024)
    if file_size_mb > MAX_FILE_SIZE_MB:
        return False, f"Template file too large: {file_size_mb:.2f}MB (max: {MAX_FILE_SIZE_MB}MB)"
    
    if not os.access(template_path, os.R_OK):
        return False, f"No read permission for file: {template_path}"
    
    try:
        doc = Document(template_path)
    except Exception as e:
        return False, f"Invalid or corrupted .docx file: {str(e)[:100]}"
    
    return True, None


def validate_variables(variables: dict) -> Tuple[bool, Optional[str]]:
    """Validate the variables dictionary."""
    if not isinstance(variables, dict):
        return False, "Variables must be a dictionary"
    
    if not variables:
        return False, "No variables provided. At least one variable is required."
    
    if len(variables) > MAX_VARIABLES_COUNT:
        return False, f"Too many variables: {len(variables)} (max: {MAX_VARIABLES_COUNT})"
    
    for key, value in variables.items():
        if not isinstance(key, str):
            return False, f"Variable name must be a string, got: {type(key).__name__}"
        
        if not key:
            return False, "Variable name cannot be empty"
        
        if not re.match(r'^[a-zA-Z_][a-zA-Z0-9_]*$', key):
            return False, f"Invalid variable name '{key}'. Must start with letter/underscore and contain only alphanumeric characters and underscores."
        
        if len(key) > 100:
            return False, f"Variable name too long: '{key}' ({len(key)} characters, max: 100)"
        
        if value is None:
            return False, f"Variable '{key}' has None value. Provide a valid value."
        
        try:
            str_value = str(value)
        except Exception as e:
            return False, f"Cannot convert variable '{key}' to string: {e}"
        
        if len(str_value) > MAX_VARIABLE_VALUE_LENGTH:
            return False, f"Variable '{key}' value too long: {len(str_value)} characters (max: {MAX_VARIABLE_VALUE_LENGTH})"
        
        forbidden_found = [char for char in FORBIDDEN_CHARS if char in str_value]
        if forbidden_found:
            return False, f"Variable '{key}' contains forbidden characters: {forbidden_found}"
    
    return True, None


def sanitize_filename(filename: str) -> str:
    """Sanitize filename to remove dangerous characters."""
    for char in FORBIDDEN_CHARS:
        filename = filename.replace(char, '_')
    
    filename = ''.join(char for char in filename if ord(char) >= 32 or char in ['\n', '\t'])
    filename = filename.strip()
    
    if not filename:
        filename = "document"
    
    if len(filename) > MAX_FILENAME_LENGTH:
        filename = filename[:MAX_FILENAME_LENGTH]
    
    return filename


# ============================================================================
# LLM SETUP
# ============================================================================

# Initialize OpenAI client
llm = LiteLlm(model="gpt-4o-mini")

def get_available_templates(directory_path: str) -> List[str]:
    templates = []
    try:
        for filename in os.listdir(directory_path):
            if filename.endswith("_template.docx"):
                friendly_name = filename.replace("_template.docx", "").replace("_", " ")
                templates.append(friendly_name)
    except Exception:
        pass
    return sorted(templates)



def find_template_with_llm(user_query: str, directory_path: str) -> dict:
    if not os.path.isdir(directory_path):
        return {"status": "error", "message": "Template directory not found."}
    
    available_templates = get_available_templates(directory_path)
    if not available_templates:
        return {"status": "error", "message": "No templates found."}

    prompt = f"""
    You are an HR assistant. Match the user request to the best template.

    User Request: "{user_query}"

    Available Templates:
    {chr(10).join(f"- {t}" for t in available_templates)}

    Respond with JSON:
    {{
        "matched_template": "exact template name",
        "confidence": 0.0 to 1.0,
        "reason": "short explanation"
    }}
    """

    # Use LiteLlm to get the response
    response = llm.generate(prompt, max_output_tokens=200)
    try:
        result = json.loads(response)
        return result
    except:
        return {"status": "no_match", "available_templates": available_templates}

def get_variables(template_path: str) -> dict:
    """Extract variables from a template."""
    is_valid, error_msg = validate_template_path(template_path)
    if not is_valid:
        return {
            "status": "error",
            "error_type": "invalid_template_path",
            "message": error_msg
        }
    
    try:
        doc = Document(template_path)
        full_text = []
        
        for para in doc.paragraphs: 
            full_text.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs: 
                        full_text.append(para.text)
        
        placeholders = re.findall(r'\{(\w+)\}', "\n".join(full_text))
        unique_placeholders = list(set(placeholders))
        
        if not unique_placeholders:
            return {
                "status": "no_variables",
                "message": "No variables found in the template.",
                "user_message": "This template doesn't contain any variables to fill."
            }
        
        return {
            "status": "variables_found",
            "variables": sorted(unique_placeholders),
            "message": f"Found {len(unique_placeholders)} variables in the template. Please provide values for each variable.",
            "variable_count": len(unique_placeholders)
        }
        
    except Exception as e: 
        return {
            "status": "error",
            "error_type": "document_read_error",
            "message": f"Error reading the .docx file: {e}",
            "user_message": "I couldn't read the template file. It may be corrupted."
        }


def fill_template(template_path: str, variables: dict) -> dict:
    """Fill the template with the provided variables and generate the document."""
    is_valid, error_msg = validate_template_path(template_path)
    if not is_valid:
        return {
            "status": "error",
            "error_type": "invalid_template_path",
            "message": error_msg,
            "user_message": "There's a problem with the template file."
        }
    
    is_valid, error_msg = validate_variables(variables)
    if not is_valid:
        return {
            "status": "error",
            "error_type": "invalid_variables",
            "message": error_msg,
            "user_message": "There's a problem with the variable values you provided."
        }
    
    try:
        doc = Document(template_path)
        replacements = {f"{{{key}}}": str(value) for key, value in variables.items()}
        replacements_made = 0
        
        # Helper function to replace placeholders in paragraphs (handles split runs)
        def replace_in_paragraph(paragraph, replacements):
            nonlocal replacements_made
            full_text = "".join(run.text for run in paragraph.runs)
            
            # Check if any placeholder exists in the paragraph
            has_replacement = False
            for placeholder, value in replacements.items():
                if placeholder in full_text:
                    has_replacement = True
                    full_text = full_text.replace(placeholder, value)
                    replacements_made += 1
            
            # If replacements were made, rebuild the paragraph
            if has_replacement:
                for run in paragraph.runs:
                    run.text = ""
                if paragraph.runs:
                    paragraph.runs[0].text = full_text
                else:
                    paragraph.add_run(full_text)
        
        # Replace in paragraphs
        for para in doc.paragraphs:
            replace_in_paragraph(para, replacements)
        
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        replace_in_paragraph(para, replacements)
        
        if replacements_made == 0:
            return {
                "status": "warning",
                "message": "Document processed but no placeholders were found to replace.",
                "user_message": "I processed the document, but didn't find any matching variables to replace. Please verify your variable names.",
                "variables_provided": list(variables.keys())
            }
        
        os.makedirs(RESULTS_DIR, exist_ok=True)
        
        base_filename = os.path.basename(template_path).replace("_template.docx", "")
        base_filename = sanitize_filename(base_filename)
        
        timestamp = int(time.time())
        unique_id = uuid.uuid4().hex[:8]
        new_filename = f"{base_filename}_{timestamp}_{unique_id}.docx"
        new_filepath = os.path.join(RESULTS_DIR, new_filename)
        
        doc.save(new_filepath)
        
        if not os.path.exists(new_filepath):
            return {
                "status": "error",
                "error_type": "save_verification_failed",
                "message": "Document save completed but file not found at expected location",
                "user_message": "Something went wrong while saving the document."
            }
        
        return {
            "status": "success",
            "message": f"Document generated successfully with {replacements_made} replacements. The file is saved as: {new_filepath}",
            "file_path": new_filepath,
            "filename": new_filename,
            "replacements_made": replacements_made,
            "variables_used": list(variables.keys())
        }
        
    except PermissionError as e:
        return {
            "status": "error",
            "error_type": "permission_denied",
            "message": f"Permission denied: {e}",
            "user_message": "I don't have permission to save the file. Please check folder permissions."
        }
    except OSError as e:
        return {
            "status": "error",
            "error_type": "io_error",
            "message": f"File system error: {e}",
            "user_message": "I encountered a file system error while saving the document."
        }
    except Exception as e:
        return {
            "status": "error",
            "error_type": "unexpected_error",
            "message": f"Unexpected error generating document: {e}",
            "user_message": "An unexpected error occurred. Please try again."
        }


# ============================================================================
# STATEFUL ASSISTANT
# ============================================================================

class DonnaAssistant:
    def __init__(self):
        self.state = "idle"
        self.template_path = None
        self.variables_needed = []
        self.variables_collected = {}

    def process_message(self, user_input: str) -> str:
        user_lower = user_input.lower().strip()
        if self.state == "idle":
            # Try to match template
            result = find_template_with_llm(user_input, TEMPLATES_DIR)
            
            if "matched_template" in result:
                # Find path
                matched_name = result["matched_template"]
                for f in os.listdir(TEMPLATES_DIR):
                    if f.endswith("_template.docx") and matched_name.lower() in f.lower():
                        self.template_path = os.path.join(TEMPLATES_DIR, f)
                        break
                
                var_result = get_variables(self.template_path)
                if var_result.get("status") == "variables_found":
                    self.variables_needed = var_result["variables"]
                    self.variables_collected = {}
                    self.state = "collecting_variables"
                    return f"Template '{matched_name}' selected. What is the {self.variables_needed[0]}?"
                else:
                    return "No variables found in the template."
            
            return "Couldn't find a matching template."

        elif self.state == "collecting_variables":
            idx = len(self.variables_collected)
            current_var = self.variables_needed[idx]

            is_valid, err, normalized = validate_input_by_type(current_var, user_input)
            if not is_valid:
                return f"❌ {err}. Please enter {current_var} again."

            self.variables_collected[current_var] = normalized or user_input.strip()
            if len(self.variables_collected) == len(self.variables_needed):
                self.state = "idle"
                result = fill_template(self.template_path, self.variables_collected)
                if result.get("status") == "success":
                    return f"✓ Document created: {result.get('file_path')}"
                else:
                    return f"Error: {result.get('user_message', 'Unknown error')}"
            else:
                next_var = self.variables_needed[idx + 1]
                return f"✓ Got it! What is the {next_var}?"

        return "I didn't understand that."

    def reset(self):
        self.state = "idle"
        self.template_path = None
        self.variables_needed = []
        self.variables_collected = {}



# ============================================================================
# MAIN ASYNC LOOP
# ============================================================================

async def run_donna_assistant():
    print("=" * 60)
    print("Donna - HR Document Assistant (Google ADK version)")
    print("=" * 60)
    print("Type 'quit' or 'exit' to end the session.\n")

    donna = DonnaAssistant()
    print("Agent: Hello! What document would you like to create today?\n")

    while True:
        try:
            user_input = await asyncio.to_thread(input, "You: ")
        except (EOFError, KeyboardInterrupt):
            print("\nGoodbye!")
            break

        if user_input.lower() in ['quit', 'exit']:
            print("\nAgent: Goodbye!")
            break

        if not user_input.strip():
            continue

        response = donna.process_message(user_input)
        print(f"\nAgent: {response}\n")

# ============================================================================
# EXECUTION
# ============================================================================

if __name__ == "__main__":
    logging.getLogger("google.adk").setLevel(logging.CRITICAL)
    try:
        asyncio.run(run_donna_assistant())
    except KeyboardInterrupt:
        print("\nGoodbye!")